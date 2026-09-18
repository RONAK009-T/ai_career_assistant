import io
from pathlib import Path
import pypdf
import docx

def extract_text_from_file(uploaded_file) -> str:
    """Extracts raw text from PDF, DOCX, or TXT/MD files."""
    if uploaded_file is None:
        return ""
    
    filename = uploaded_file.name
    ext = Path(filename).suffix.lower()
    
    try:
        if ext == ".pdf":
            reader = pypdf.PdfReader(io.BytesIO(uploaded_file.read()))
            text_parts = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text_parts.append(extracted)
            return "\n\n".join(text_parts).strip()
            
        elif ext in [".docx", ".doc"]:
            doc = docx.Document(io.BytesIO(uploaded_file.read()))
            text_parts = [p.text for p in doc.paragraphs if p.text]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        text_parts.append(row_text)
            return "\n".join(text_parts).strip()
            
        elif ext in [".txt", ".md"]:
            raw = uploaded_file.read()
            try:
                return raw.decode("utf-8").strip()
            except UnicodeDecodeError:
                return raw.decode("latin-1", errors="ignore").strip()
                
        else:
            return ""
    except Exception as e:
        print(f"Error parsing file {filename}: {e}")
        return ""
